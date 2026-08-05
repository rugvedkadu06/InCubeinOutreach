import os
import re
import json
from datetime import datetime

from ..core import config
from ..core.database import get_db_connection
from ..core.exceptions import NotFoundError, ServiceError
from ..schemas.incubators import ContactUpdateRequest, FinderRequest


def _find_excel_path():
    excel_path = config.APP_DATA_DIR / "incubators_with_contact_details(100).xlsx"
    if not excel_path.exists():
        excel_path = config.APP_DATA_DIR / "db01.xlsx"
    return excel_path


def update_incubator_contact(req: ContactUpdateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()

    # 1. Fetch current details to get the canonical name of the incubator
    cursor.execute("SELECT name FROM incubators WHERE id = ?", (req.id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise NotFoundError("Incubator not found")

    inc_name = row["name"]

    # 2. Update the SQLite database
    try:
        cursor.execute(
            "UPDATE incubators SET email = ?, website = ?, last_updated = ? WHERE id = ?",
            (req.email.strip(), req.website.strip(), datetime.now().isoformat(), req.id)
        )
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        raise ServiceError(f"Database update failed: {str(e)}")

    conn.close()

    # 3. Update the Excel file (incubators_with_contact_details(100).xlsx or db01.xlsx)
    excel_updated = False
    excel_error = None
    try:
        import openpyxl
        excel_path = _find_excel_path()
        if os.path.exists(str(excel_path)):
            wb = openpyxl.load_workbook(str(excel_path))
            sheet = wb.active

            rows_updated = 0
            for r in range(2, sheet.max_row + 1):
                cell_val = sheet.cell(row=r, column=1).value
                if cell_val:
                    # Check exact match or similarity
                    if str(cell_val).strip().lower() == inc_name.strip().lower():
                        sheet.cell(row=r, column=4).value = req.website.strip()  # Column D (Website)
                        sheet.cell(row=r, column=5).value = req.email.strip()    # Column E (Email)
                        rows_updated += 1

            # If no exact match, find the best match using similarity
            if rows_updated == 0:
                best_row = None
                best_sim = 0.0

                # Jaccard token overlap similarity helper
                def get_words(t):
                    return {w for w in re.sub(r'[^a-zA-Z0-9\s]', ' ', t.lower()).split() if w not in {"for", "and", "the", "of", "in", "at", "on", "with", "a", "an"}}

                name_words = get_words(inc_name)
                for r in range(2, sheet.max_row + 1):
                    cell_val = sheet.cell(row=r, column=1).value
                    if cell_val:
                        k_words = get_words(str(cell_val))
                        if name_words and k_words:
                            sim = len(name_words.intersection(k_words)) / len(name_words.union(k_words))
                            if str(cell_val).lower() in inc_name.lower() or inc_name.lower() in str(cell_val).lower():
                                sim = max(sim, 0.6)
                            if sim > best_sim:
                                best_sim = sim
                                best_row = r

                if best_row and best_sim >= 0.4:
                    sheet.cell(row=best_row, column=4).value = req.website.strip()
                    sheet.cell(row=best_row, column=5).value = req.email.strip()
                    rows_updated = 1

            if rows_updated > 0:
                wb.save(str(excel_path))
                excel_updated = True
            else:
                excel_error = f"Could not find matching incubator row in {excel_path.name}"
        else:
            excel_error = f"{excel_path.name} file not found"
    except Exception as e:
        excel_error = f"Failed to write to {excel_path.name}: {str(e)}"

    return {
        "status": "success",
        "message": "Incubator contact details updated in SQLite.",
        "excel_status": "updated" if excel_updated else "skipped",
        "excel_error": excel_error
    }


def find_matching_incubators(req: FinderRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM incubators")
    incubators = [dict(row) for row in cursor.fetchall()]
    conn.close()

    matches = []
    startup_sector = req.sector.strip().lower()
    startup_city = req.hq_city.strip().lower()
    startup_state = req.state.strip().lower() if req.state else ""
    startup_stage = req.stage.strip().lower()

    for inc in incubators:
        score = 0
        reasons = []

        # 1. Sector matching (40 points)
        focus_areas = []
        if inc["focus_areas"]:
            try:
                focus_areas = json.loads(inc["focus_areas"])
            except:
                focus_areas = []

        focus_areas_lower = [f.lower() for f in focus_areas]
        if startup_sector in focus_areas_lower:
            score += 40
            reasons.append(f"Direct Sector Match: Incubator specializes in {req.sector}")
        elif any(startup_sector in f or f in startup_sector for f in focus_areas_lower):
            score += 30
            reasons.append(f"Sub-Sector Match: Specializes in related areas ({', '.join(focus_areas[:3])})")
        elif not focus_areas:
            score += 15
            reasons.append("General Sector Support: Multi-sector incubator open to all verticals")
        else:
            score += 5

        # 2. Location matching (30 points)
        inc_city = (inc["city"] or "").strip().lower()
        inc_state = (inc["state"] or "").strip().lower()
        if startup_city and inc_city == startup_city:
            score += 30
            reasons.append(f"Local City Hub: Located in {inc['city']}, matching your startup HQ")
        elif startup_state and inc_state == startup_state:
            score += 20
            reasons.append(f"Regional State Hub: Located in {inc['city']}, {inc['state']}")
        else:
            score += 8
            reasons.append(f"National Incubator: Located in {inc['city'] or 'India'}")

        # 3. Funding & Scheme matching (20 points)
        grant_support = (inc["grant_support"] or "").strip().lower()
        funding_support = (inc["funding_support"] or "").strip().lower()

        has_grant = any(w in grant_support or w in funding_support for w in ["prayas", "birac", "big", "sisfs", "nidhi", "seed", "grant"])
        if startup_stage in ["pre-seed", "seed", "idea"]:
            if has_grant:
                score += 20
                reasons.append(f"Early-Stage Support: Offers seed grants/schemes like {inc['grant_support'] or 'NIDHI/BIRAC'}")
            else:
                score += 12
                reasons.append("Early-Stage Match: Mentorship & incubation space available")
        else:
            score += 10
            reasons.append("Growth Support: Coworking and networking infrastructure")

        # 4. Confidence/Capability matching (10 points)
        conf = inc["confidence_score"] or 0.90
        score += int(conf * 10)

        score = min(100, score)

        # Clean incubator JSON fields for output
        for field in ["incubation_programs", "acceleration_programs", "lab_facilities", "focus_areas"]:
            if inc[field] and isinstance(inc[field], str):
                try:
                    inc[field] = json.loads(inc[field])
                except:
                    inc[field] = []

        matches.append({
            "incubator": inc,
            "match_score": score,
            "reasons": reasons
        })

    matches.sort(key=lambda x: x["match_score"], reverse=True)
    return matches[:5]
